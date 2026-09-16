"""Encode a set of exported rows with the real formatters and print what they produce.

The Go side encodes the same rows and diffs the two, which is what keeps the csv, the json and
the spreadsheet a port rather than a guess. The xlsx bytes are not comparable — two writers never
produce the same zip — so what is recorded for it is the cell grid openpyxl reads back out.

Run from apps/api with the project's settings on the path:

    DJANGO_SETTINGS_MODULE=plane.settings.production SECRET_KEY=x \
    REDIS_URL=redis://localhost:6379/ DATABASE_URL=postgresql://u:p@localhost:5432/plane \
    python ../api-go/tools/generate_export_format_fixture.py > \
    ../api-go/internal/worker/testdata/export_format.json
"""

import json
import os
import sys

sys.path.insert(0, os.getcwd())

import django  # noqa: E402

django.setup()

from plane.utils.porters.formatters import CSVFormatter, JSONFormatter, XLSXFormatter  # noqa: E402

# One ordinary row, one with everything empty, and one made entirely of the things that break an
# encoder: formula triggers, quotes, newlines, non-ASCII text and an apostrophe inside an object.
ROWS = [
    {
        "project_name": "Apollo",
        "project_identifier": "APO",
        "parent": "APO-1",
        "identifier": "APO-42",
        "sequence_id": 42,
        "name": "Land the thing",
        "state_name": "In Progress",
        "priority": "urgent",
        "assignees": ["Ada Lovelace", "Grace Hopper"],
        "subscribers": ["Ada Lovelace"],
        "created_by_name": "Ada Lovelace",
        "start_date": "2026-01-01",
        "target_date": "2026-02-01",
        "completed_at": None,
        "created_at": "2026-01-01T00:00:00Z",
        "updated_at": "2026-01-02T03:04:05.120000Z",
        "archived_at": None,
        "estimate": "8",
        "labels": ["backend", "urgent"],
        "cycles": ["Sprint 1"],
        "modules": ["Launch"],
        "links": [{"url": "https://example.test/a", "title": "A link"}],
        "relations": [{"type": "blocked_by", "issue": "APO-7", "direction": "outgoing"}],
        "comments": [
            {"comment": "ship it", "created_by": "Ada Lovelace", "created_at": "2026-01-01 10:00:00"}
        ],
        "is_draft": False,
    },
    {
        "project_name": "Apollo",
        "project_identifier": "APO",
        "parent": "",
        "identifier": "APO-43",
        "sequence_id": 43,
        "name": "",
        "state_name": "",
        "priority": "none",
        "assignees": [],
        "subscribers": [],
        "created_by_name": "",
        "start_date": None,
        "target_date": None,
        "completed_at": None,
        "created_at": "2026-01-03T00:00:00Z",
        "updated_at": "2026-01-03T00:00:00Z",
        "archived_at": None,
        "estimate": "",
        "labels": [],
        "cycles": [],
        "modules": [],
        "links": [],
        "relations": [],
        "comments": [],
        "is_draft": True,
    },
    {
        "project_name": "=cmd|' /C calc'!A0",
        "project_identifier": "-BAD",
        "parent": "+1",
        "identifier": "@here",
        "sequence_id": 44,
        "name": 'a "quoted", comma\nand a newline',
        "state_name": "Ünicode ☃",
        "priority": "low",
        "assignees": ["O'Brien", "Zoë"],
        "subscribers": [],
        "created_by_name": "\tTabbed",
        "start_date": None,
        "target_date": None,
        "completed_at": "2026-01-04T05:06:07Z",
        "created_at": "2026-01-04T00:00:00Z",
        "updated_at": "2026-01-04T00:00:00Z",
        "archived_at": "2026-01-05",
        "estimate": "",
        "labels": ["back\\slash"],
        "cycles": [],
        "modules": [],
        "links": [{"url": "https://example.test/b", "title": "it's here"}],
        "relations": [],
        "comments": [
            {"comment": "line one\nline two", "created_by": "Zoë", "created_at": "2026-01-04 01:02:03"}
        ],
        "is_draft": False,
    },
]


def xlsx_grid(rows):
    """Encode with openpyxl and read the cells straight back, since the bytes are not comparable."""
    from io import BytesIO

    from openpyxl import load_workbook

    content = XLSXFormatter(list_joiner=", ").encode(rows)
    sheet = load_workbook(filename=BytesIO(content), data_only=True).active
    return {
        "title": sheet.title,
        "cells": [[cell for cell in row] for row in sheet.iter_rows(values_only=True)],
    }


def case(name, rows):
    return {
        "name": name,
        "rows": rows,
        "csv": CSVFormatter().encode(rows),
        "json": JSONFormatter().encode(rows),
        "xlsx": xlsx_grid(rows),
    }


print(
    json.dumps(
        [case("rows", ROWS), case("empty", [])],
        indent=2,
        ensure_ascii=False,
    )
)
